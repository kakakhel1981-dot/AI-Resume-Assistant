import io
import json
import os
import re
from pathlib import Path
from typing import Any

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# ------------------------------------------------------------
# Page configuration
# ------------------------------------------------------------
st.set_page_config(
    page_title="Resume ATS Analyzer",
    page_icon="📄",
    layout="wide",
)

# Gemini Flash model. You can change this in Streamlit Secrets
# with GEMINI_MODEL if you want to use another supported model.
DEFAULT_MODEL = "gemini-3.5-flash"

MAX_RESUME_CHARS = 50000
MAX_JOB_DESCRIPTION_CHARS = 30000

# Weighted scoring model. The final score is calculated locally
# from the category scores returned by Gemini.
WEIGHTS = {
    "ats_format": 20,
    "sections": 15,
    "keywords": 20,
    "experience_bullets": 20,
    "skills": 10,
    "contact": 5,
    "education": 5,
    "readability": 5,
}


# ------------------------------------------------------------
# Configuration helpers
# ------------------------------------------------------------
def get_secret_or_env(name: str, default: str | None = None) -> str | None:
    """Read a value from Streamlit Secrets first, then environment variables."""
    try:
        value = st.secrets.get(name)
        if value:
            return str(value).strip()
    except Exception:
        pass

    value = os.getenv(name)
    return value.strip() if value else default


def get_model_name() -> str:
    return get_secret_or_env("GEMINI_MODEL", DEFAULT_MODEL) or DEFAULT_MODEL


def get_api_key() -> str | None:
    return get_secret_or_env("GEMINI_API_KEY")


# ------------------------------------------------------------
# Resume text extraction
# ------------------------------------------------------------
def extract_pdf_text(file_bytes: bytes) -> str:
    """Extract text from a normal text-based PDF."""
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = []

    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")

    return "\n".join(pages).strip()


def extract_docx_text(file_bytes: bytes) -> str:
    """Extract paragraphs and table content from a DOCX file."""
    document = Document(io.BytesIO(file_bytes))
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()


def extract_text_file(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="replace").strip()


def extract_resume_text(uploaded_file) -> str:
    """Extract readable text based on the uploaded file extension."""
    file_bytes = uploaded_file.getvalue()
    extension = Path(uploaded_file.name).suffix.lower()

    if extension == ".pdf":
        text = extract_pdf_text(file_bytes)
    elif extension == ".docx":
        text = extract_docx_text(file_bytes)
    elif extension in {".txt", ".md"}:
        text = extract_text_file(file_bytes)
    else:
        raise ValueError(
            "Unsupported file type. Please upload PDF, DOCX, TXT, or MD."
        )

    if not text.strip():
        raise ValueError(
            "No readable text was extracted. If your PDF is scanned or "
            "image-only, please upload an OCR-enabled PDF or DOCX file."
        )

    return text[:MAX_RESUME_CHARS]


# ------------------------------------------------------------
# Gemini prompt
# ------------------------------------------------------------
def build_analysis_prompt(resume_text: str, job_description: str) -> str:
    job_description = job_description.strip()

    if job_description:
        job_section = (
            "JOB DESCRIPTION:\n"
            + job_description[:MAX_JOB_DESCRIPTION_CHARS]
        )
        keyword_instruction = (
            "Compare the resume against the job description. Identify keywords "
            "that are clearly present, missing, or weakly evidenced."
        )
    else:
        job_section = (
            "JOB DESCRIPTION:\n"
            "Not provided. Perform a general ATS-readiness assessment."
        )
        keyword_instruction = (
            "Because no job description was provided, evaluate general "
            "role-relevant terminology and standard professional skills."
        )

    return f"""
You are an expert ATS resume reviewer, recruiter, and resume-writing specialist.

Analyze the supplied resume for ATS-readiness and recruiter readability.

IMPORTANT RULES:
1. Base every finding only on the supplied resume and job description.
2. Never invent employers, job titles, dates, degrees, certifications, skills,
   technologies, achievements, metrics, or responsibilities.
3. Never recommend adding a keyword unless the candidate genuinely has that
   skill or experience.
4. Do not recommend keyword stuffing.
5. Focus on practical ATS concerns such as standard headings, readable text,
   unusual formatting, columns/tables, graphics, icons, headers/footers,
   inconsistent dates, and unclear sections.
6. {keyword_instruction}
7. Bullet rewrites may improve grammar, clarity, action verbs, and impact, but
   must use only facts already present in the resume.
8. Keep recommendations specific and actionable.
9. Return ONLY valid JSON. Do not return Markdown or explanatory text outside JSON.

SCORING CATEGORIES:
Score each category from 0 to 100.

- ats_format:
  Parser-friendly formatting, standard headings, readable structure, and
  absence of obvious ATS parsing risks.

- sections:
  Presence and clarity of appropriate sections such as Summary, Experience,
  Education, Skills, Certifications, Projects, etc.

- keywords:
  Job-description alignment when a JD is provided. Otherwise, quality of
  standard role-relevant terminology.

- experience_bullets:
  Action verbs, specificity, measurable outcomes, concise writing, and
  evidence of impact rather than vague responsibilities.

- skills:
  Clear and relevant technical/professional skills.

- contact:
  Clear professional contact information such as name, email, phone,
  LinkedIn/portfolio where appropriate.

- education:
  Clear and appropriately presented education information.

- readability:
  Grammar, consistency, concise wording, date consistency, and scanability.

Return exactly this JSON structure:

{{
  "summary": "Overall assessment in 3-5 sentences.",
  "category_scores": {{
    "ats_format": 0,
    "sections": 0,
    "keywords": 0,
    "experience_bullets": 0,
    "skills": 0,
    "contact": 0,
    "education": 0,
    "readability": 0
  }},
  "category_evidence": {{
    "ats_format": "Evidence supporting the score.",
    "sections": "Evidence supporting the score.",
    "keywords": "Evidence supporting the score.",
    "experience_bullets": "Evidence supporting the score.",
    "skills": "Evidence supporting the score.",
    "contact": "Evidence supporting the score.",
    "education": "Evidence supporting the score.",
    "readability": "Evidence supporting the score."
  }},
  "strengths": [
    "Strength 1",
    "Strength 2",
    "Strength 3"
  ],
  "improvements": [
    {{
      "priority": "High",
      "issue": "Specific issue",
      "recommendation": "Specific recommendation"
    }}
  ],
  "keyword_analysis": {{
    "present": ["keyword1", "keyword2"],
    "missing_or_weak": ["keyword3", "keyword4"],
    "notes": "Additional keyword observations."
  }},
  "bullet_rewrites": [
    {{
      "original": "Original resume bullet",
      "improved": "Improved version using only existing facts",
      "reason": "Why the revised bullet is stronger"
    }}
  ],
  "ats_checklist": [
    "Checklist item 1",
    "Checklist item 2"
  ]
}}

RESUME:
{resume_text}

{job_section}
"""


# ------------------------------------------------------------
# Scoring / JSON normalization
# ------------------------------------------------------------
def normalize_result(result: dict[str, Any]) -> dict[str, Any]:
    """Validate model output and calculate the weighted ATS score locally."""
    scores = result.get("category_scores", {})
    normalized_scores: dict[str, int] = {}

    for category in WEIGHTS:
        try:
            score = int(round(float(scores.get(category, 0))))
        except (TypeError, ValueError):
            score = 0

        normalized_scores[category] = max(0, min(100, score))

    result["category_scores"] = normalized_scores

    weighted_score = sum(
        normalized_scores[category] * weight
        for category, weight in WEIGHTS.items()
    ) / 100

    result["ats_score"] = round(weighted_score)

    if not isinstance(result.get("strengths"), list):
        result["strengths"] = []

    if not isinstance(result.get("improvements"), list):
        result["improvements"] = []

    if not isinstance(result.get("bullet_rewrites"), list):
        result["bullet_rewrites"] = []

    if not isinstance(result.get("ats_checklist"), list):
        result["ats_checklist"] = []

    if not isinstance(result.get("category_evidence"), dict):
        result["category_evidence"] = {}

    if not isinstance(result.get("keyword_analysis"), dict):
        result["keyword_analysis"] = {}

    keyword_data = result["keyword_analysis"]

    if not isinstance(keyword_data.get("present"), list):
        keyword_data["present"] = []

    if not isinstance(keyword_data.get("missing_or_weak"), list):
        keyword_data["missing_or_weak"] = []

    if not isinstance(keyword_data.get("notes"), str):
        keyword_data["notes"] = ""

    return result


# ------------------------------------------------------------
# Gemini analysis
# ------------------------------------------------------------
def analyze_resume(
    resume_text: str,
    job_description: str,
) -> dict[str, Any]:
    api_key = get_api_key()

    if not api_key:
        raise RuntimeError(
            "GEMINI_API_KEY was not found. Add it to Streamlit Secrets "
            "or set it as an environment variable."
        )

    model_name = get_model_name()
    client = genai.Client(api_key=api_key)

    response = client.models.generate_content(
        model=model_name,
        contents=build_analysis_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )

    raw_text = (response.text or "").strip()

    if not raw_text:
        raise RuntimeError(
            "Gemini returned an empty response. Please try again."
        )

    # Defensive cleanup if the model accidentally adds a JSON code fence.
    raw_text = re.sub(
        r"^```(?:json)?\s*",
        "",
        raw_text,
        flags=re.IGNORECASE,
    )
    raw_text = re.sub(r"\s*```$",
                      "",
                      raw_text)

    try:
        result = json.loads(raw_text)
    except json.JSONDecodeError as exc:
        raise RuntimeError(
            "Gemini returned invalid JSON. Please click Analyze again."
        ) from exc

    if not isinstance(result, dict):
        raise RuntimeError("Gemini returned an unexpected response format.")

    return normalize_result(result)


# ------------------------------------------------------------
# UI helpers
# ------------------------------------------------------------
def get_score_message(score: int) -> tuple[str, str]:
    if score >= 85:
        return "Excellent ATS readiness", "success"
    if score >= 70:
        return "Good ATS readiness with room for improvement", "info"
    if score >= 55:
        return "Several improvements are recommended", "warning"
    return "Significant ATS improvements are recommended", "error"


def render_results(result: dict[str, Any]) -> None:
    score = int(result["ats_score"])

    st.divider()
    st.subheader("🎯 Estimated ATS Readiness Score")

    score_col, message_col = st.columns([1, 3])

    with score_col:
        st.metric("ATS Score", f"{score}/100")

    with message_col:
        message, level = get_score_message(score)
        if level == "success":
            st.success(message)
        elif level == "info":
            st.info(message)
        elif level == "warning":
            st.warning(message)
        else:
            st.error(message)

    st.progress(score / 100)

    st.caption(
        "Important: this is an AI-based ATS-readiness estimate, not the actual "
        "score from a specific employer's ATS. Different ATS platforms, "
        "employers, and job descriptions may produce different outcomes."
    )

    # Score breakdown
    st.subheader("📊 Score Breakdown")

    categories = list(WEIGHTS.keys())

    for start in range(0, len(categories), 4):
        columns = st.columns(4)

        for column, category in zip(
            columns,
            categories[start:start + 4],
        ):
            with column:
                st.metric(
                    category.replace("_", " ").title(),
                    f"{result['category_scores'][category]}/100",
                )

    with st.expander("Why did I receive these scores?"):
        evidence = result.get("category_evidence", {})

        for category in categories:
            st.markdown(
                f"**{category.replace('_', ' ').title()} — "
                f"{result['category_scores'][category]}/100**"
            )
            st.write(
                evidence.get(
                    category,
                    "No explanation was returned.",
                )
            )

    # Summary
    st.subheader("📝 Overall Assessment")
    st.write(result.get("summary", "No summary was returned."))

    # Strengths and improvements
    left, right = st.columns(2)

    with left:
        st.markdown("### ✅ Strengths")

        strengths = result.get("strengths", [])

        if strengths:
            for strength in strengths:
                st.success(str(strength))
        else:
            st.write("No strengths were returned.")

    with right:
        st.markdown("### 🔧 Priority Improvements")

        improvements = result.get("improvements", [])

        if improvements:
            for item in improvements:
                if isinstance(item, dict):
                    priority = item.get("priority", "Medium")
                    issue = item.get("issue", "")
                    recommendation = item.get("recommendation", "")

                    st.markdown(f"**{priority}: {issue}**")
                    st.write(recommendation)
                else:
                    st.write(str(item))
        else:
            st.write("No improvements were returned.")

    # Keyword analysis
    st.subheader("🔑 Keyword Analysis")

    keyword_data = result.get("keyword_analysis", {})

    keyword_col_1, keyword_col_2 = st.columns(2)

    with keyword_col_1:
        st.markdown("**Present / Clearly Evidenced**")

        present = keyword_data.get("present", [])

        if present:
            st.write(", ".join(str(x) for x in present))
        else:
            st.write("No specific keywords were identified.")

    with keyword_col_2:
        st.markdown("**Missing or Weak**")

        missing = keyword_data.get("missing_or_weak", [])

        if missing:
            st.write(", ".join(str(x) for x in missing))
        else:
            st.write("No major missing/weak keywords were identified.")

    notes = keyword_data.get("notes", "")

    if notes:
        st.caption(notes)

    # Bullet rewrites
    st.subheader("✍️ Resume Bullet Improvements")

    rewrites = result.get("bullet_rewrites", [])

    if not rewrites:
        st.write("No bullet rewrites were suggested.")
    else:
        for index, item in enumerate(rewrites, start=1):
            if not isinstance(item, dict):
                continue

            with st.expander(f"Bullet Rewrite {index}"):
                st.markdown("**Original**")
                st.write(item.get("original", ""))

                st.markdown("**Improved**")
                st.write(item.get("improved", ""))

                st.caption(item.get("reason", ""))

    # Checklist
    st.subheader("☑️ ATS Checklist")

    checklist = result.get("ats_checklist", [])

    if checklist:
        for item in checklist:
            st.write(f"☐ {item}")
    else:
        st.write("No checklist items were returned.")

    # Download analysis JSON
    st.subheader("⬇️ Export Analysis")

    json_data = json.dumps(
        result,
        indent=2,
        ensure_ascii=False,
    )

    st.download_button(
        label="Download Analysis as JSON",
        data=json_data,
        file_name="resume_ats_analysis.json",
        mime="application/json",
    )


# ------------------------------------------------------------
# Main application
# ------------------------------------------------------------
st.title("📄 Resume ATS Analyzer")

st.write(
    "Upload a resume and optionally paste a target job description. "
    "The app uses Gemini Flash to analyze ATS-readiness, keywords, "
    "resume structure, skills, and experience bullets."
)

with st.sidebar:
    st.header("⚙️ How It Works")

    st.markdown(
        """
**1. Upload Resume**

Supported formats:
- PDF
- DOCX
- TXT
- MD

**2. Add Job Description**

Optional, but recommended for job-specific keyword analysis.

**3. Analyze**

Gemini reviews the resume.

**4. Improve**

Get:
- ATS score
- Score breakdown
- Strengths
- Priority improvements
- Keyword analysis
- Bullet rewrites
- ATS checklist
"""
    )

    st.divider()

    st.caption(f"Gemini model: {get_model_name()}")
    st.caption("Resume files are processed in memory by this application.")


resume_file = st.file_uploader(
    "📎 Upload your resume",
    type=["pdf", "docx", "txt", "md"],
    help="Upload a text-based PDF, DOCX, TXT, or MD resume.",
)

job_description = st.text_area(
    "💼 Target Job Description (Optional)",
    height=250,
    placeholder=(
        "Paste the complete job description here. "
        "For example: Frontend Developer — React, JavaScript, "
        "TypeScript, HTML, CSS, REST APIs, Git..."
    ),
)

analyze_button = st.button(
    "🚀 Analyze Resume",
    type="primary",
    disabled=resume_file is None,
    use_container_width=True,
)

if analyze_button and resume_file is not None:
    try:
        with st.spinner("📄 Extracting resume text..."):
            resume_text = extract_resume_text(resume_file)

        if len(resume_text.strip()) < 100:
            st.warning(
                "Only a small amount of text was extracted. "
                "Please verify that the resume is readable and is not "
                "an image-only/scanned PDF."
            )

        with st.spinner("🤖 Gemini is analyzing your resume..."):
            analysis = analyze_resume(
                resume_text=resume_text,
                job_description=job_description,
            )

        st.session_state["analysis"] = analysis
        st.session_state["resume_filename"] = resume_file.name

    except Exception as exc:
        st.error(f"Analysis failed: {exc}")

if "analysis" in st.session_state:
    st.caption(
        f"Analyzed file: {st.session_state.get('resume_filename', 'resume')}"
    )

    render_results(st.session_state["analysis"])
